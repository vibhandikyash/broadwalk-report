from __future__ import annotations

from pathlib import Path

import openpyxl

from .document import Document, Sheet


def read_xlsx(path: Path, file_id: str, filename: str) -> Document:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets: list[Sheet] = []
    try:
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            while rows and all(v is None for v in rows[-1]):
                rows.pop()
            sheets.append(Sheet(name=ws.title, rows=rows))
    finally:
        wb.close()
    return Document(file_id=file_id, filename=filename, kind="xlsx", sheets=sheets)
