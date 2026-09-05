# backend/tests/test_dataset.py
import os
from pathlib import Path

import pytest

DATASET = os.getenv("TEST_DATASET_DIR")
pytestmark = pytest.mark.skipif(not DATASET, reason="set TEST_DATASET_DIR to the folder holding the real source files")


def test_real_dataset_end_to_end():
    from fastapi.testclient import TestClient

    from app.main import app
    from app.workers.pool import pool

    from tests.helpers import dataset_files

    paths = dataset_files(Path(DATASET))
    with TestClient(app) as client:
        pid = client.post("/api/projects", json={"name": "dataset"}).json()["id"]
        for p in paths:
            with p.open("rb") as fh:
                assert client.post(f"/api/projects/{pid}/files", files=[("files", (p.name, fh, "application/octet-stream"))]).status_code == 201
        assert pool.wait_idle(180)
        detail = client.get(f"/api/projects/{pid}").json()
        bad = [(f["original_filename"], f["status"], f["error"]) for f in detail["files"] if f["status"] not in ("processed", "unsupported")]
        assert not bad, bad
        ui = client.get(f"/api/projects/{pid}/report-data").json()
        v = {f["path"]: f["effective"] for s in ui["sections"] for f in s["fields"]}
        cells = {c["path"]: c["effective"] for s in ui["sections"] for t in s["tables"] for r in t["rows"] for c in r["cells"]}
        cells.update({c["path"]: c["effective"] for s in ui["sections"] for t in s["tables"] for c in t["totals"]})
        assert v["property.fields.name"] == "The Boardwalk" and v["property.fields.units"] == 338
        assert v["property.fields.year_built"] == 1973 and v["property.fields.acquired_date"] == "2025-07-30"
        assert v["capital.fields.purchase_price"] == 48000000 and abs(v["capital.fields.equity_invested"] - 14259605.94) < 1
        assert abs(v["financing.fields.implied_rate"] - 0.0523) < 1e-4 and v["financing.fields.loan_amount"] == 36519000
        assert abs(cells["financials.tables.lines.rows.noi.ptd_actual"] - 636105.69) < 1
        assert abs(cells["financials.tables.lines.rows.total_revenue.ytd_actual"] - 2826719.16) < 1
        assert abs(cells["financials.tables.lines.rows.concessions.ptd_actual"] + 53131.05) < 1
        assert abs(cells["financials.tables.lines.rows.net_cash_flow.ptd_var"] + 134365.64) < 1
        assert abs(cells["capex.tables.lines.totals.ptd_actual"] - 241077.37) < 1 and abs(v["capex.fields.source_total_ptd"] - 241077.37) < 1
        assert abs(cells["capex.tables.lines.rows.plumbing_water_heaters.ptd_actual"] - 80252.75) < 1
        assert abs(cells["capex.tables.lines.totals.ptd_budget"] - 278896) < 1 and abs(cells["capex.tables.lines.totals.ytd_budget"] - 616008) < 1
        assert abs(cells["capex.tables.lines.totals.annual_budget"] - 1106416) < 1
        assert abs(cells["capex.tables.lines.rows.renovation_plumbing.ptd_budget"] - 54078) < 1
        assert abs(cells["in_place_rent.tables.by_floor_plan.totals.current_rent"] - 1323.98) < 0.01
        assert abs(cells["in_place_rent.tables.by_floor_plan.rows.0br.current_rent"] - 1074.1) < 0.5
        assert abs(v["occupancy.fields.current_pct"] - 0.9053) < 1e-4 and v["occupancy.fields.change_bps"] == -118
        assert cells["occupancy.tables.new_leases.totals.count"] == 38 and abs(cells["occupancy.tables.new_leases.totals.lto_pct"] + 0.1762) < 1e-3
        assert cells["occupancy.tables.renewals.totals.count"] == 25
        assert abs(v["submarket.fields.vacancy"] - 0.1634) < 1e-3 and abs(v["submarket.fields.prior_vacancy"] - 0.1471) < 1e-3
        assert abs(cells["submarket.tables.comps.rows.the-ashlar.asking_rent"] - 1719) < 1
        assert abs(cells["submarket.tables.comps.rows.the-boardwalk.asking_rent"] - 1303) < 1
        assert v["submarket.fields.submarket_name"] == "Western Lee County" and v["property.fields.prepared_by"] == "ZMR Capital"
        assert not [i for i in ui["issues"] if i["severity"] == "error"], [i["message"] for i in ui["issues"] if i["severity"] == "error"]


def _swap_cols(ws, a: int, b: int) -> None:
    from openpyxl.cell.cell import MergedCell

    for row in ws.iter_rows():
        if max(a, b) <= len(row) and not isinstance(row[a - 1], MergedCell) and not isinstance(row[b - 1], MergedCell):
            row[a - 1].value, row[b - 1].value = row[b - 1].value, row[a - 1].value


def _find_col(ws, pattern: str) -> int | None:
    import re

    for row in ws.iter_rows(max_row=12):
        for cell in row:
            if isinstance(cell.value, str) and re.search(pattern, cell.value, re.I):
                return cell.column
    return None


def _mutated_copy(src: Path, dst: Path) -> list[str]:
    """A second dataset built from the first: random file names, renamed sheets, inserted rows and columns,
    reordered and dropped columns, extra irrelevant sheets, alternate date and percent spellings, a duplicated
    same-period export, an unrelated property's rent roll and an image-only PDF. Returns the names of the
    files that are expected NOT to end up 'processed'."""
    import shutil
    import uuid

    import openpyxl
    import pypdfium2 as pdfium

    from tests.helpers import make_xlsx, rent_roll_rows

    dst.mkdir(parents=True, exist_ok=True)
    copies: dict[str, Path] = {}
    from tests.helpers import dataset_files

    for p in dataset_files(src):
        if "_Misc" in p.parts:
            continue
        q = dst / f"{uuid.uuid4().hex[:10]}{p.suffix.lower()}"
        shutil.copy(p, q)
        copies[p.name] = q

    def financials(ws):
        ws.insert_rows(1, amount=2)
        ws.insert_cols(1, amount=1)
        a, b = _find_col(ws, r"^\s*ptd actual"), _find_col(ws, r"^\s*ptd budget")
        _swap_cols(ws, a, b)  # header-driven, so reordering must not matter
        ws.delete_cols(_find_col(ws, r"^\s*annual"))  # an optional column that may be missing

    def rent_roll(ws):
        ws.insert_cols(1, amount=1)
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip().lower().startswith("as of"):
                    cell.value = "As Of = 2026-06-30"  # ISO instead of m/d/Y
            label = next((c.value for c in row if isinstance(c.value, str) and c.value.strip()), "")
            if str(label).strip().lower() == "occupied units":
                for cell in row:
                    if isinstance(cell.value, (int, float)) and abs(cell.value - 90.53) < 0.01:
                        cell.value = "90.53%"  # percent as text

    mutations = {
        "PTD and YTD Financials": financials,
        "Occupancy_6-30-26": rent_roll,
        "REnt Schedule_06.30.26": lambda ws: ws.insert_rows(1, amount=1),
        "Costar Submarket Excel": lambda ws: ws.insert_rows(1, amount=1),
        "Yardi LTO": lambda ws: ws.insert_cols(1, amount=1),
        "Comp Set_2Q26 Leasing": lambda ws: ws.insert_rows(1, amount=3),
    }
    for name, path in copies.items():
        fn = next((f for key, f in mutations.items() if key in name), None)
        if fn is None:
            continue
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        ws.title = "Sheet1"
        fn(ws)
        junk = wb.create_sheet("Notes")  # an extra sheet that means nothing
        junk.append(["Prepared for review", "internal"])
        junk.append(["Reviewer", "n/a", 42])
        wb.save(path)
    # the same-period financials twice (a re-export), an unrelated property's rent roll, and a scanned PDF
    fin = next(p for n, p in copies.items() if "PTD and YTD Financials" in n)
    shutil.copy(fin, dst / f"{uuid.uuid4().hex[:10]}.xlsx")
    rows = [[("Lakeside Villas (99001)" if c == "The Boardwalk (45726)" else ("Lakeside Villas" if c == "The Boardwalk" else c)) for c in r]
            for r in rent_roll_rows("06/30/2026", 6, 12, 50.0, 0)]
    make_xlsx(dst / f"{uuid.uuid4().hex[:10]}.xlsx", {"Report1": rows})
    pdf = pdfium.PdfDocument.new()
    pdf.new_page(400, 300)
    scan = dst / f"{uuid.uuid4().hex[:10]}.pdf"
    pdf.save(scan)
    return [scan.name]


def test_mutated_dataset_reproduces_key_figures(tmp_path):
    from fastapi.testclient import TestClient

    from app.main import app
    from app.workers.pool import pool

    ds2 = tmp_path / "ds2"
    not_processed = _mutated_copy(Path(DATASET), ds2)
    paths = sorted(ds2.iterdir())
    with TestClient(app) as client:
        pid = client.post("/api/projects", json={"name": "mutated"}).json()["id"]
        for p in paths:
            with p.open("rb") as fh:
                assert client.post(f"/api/projects/{pid}/files", files=[("files", (p.name, fh, "application/octet-stream"))]).status_code == 201
        assert pool.wait_idle(180)
        detail = client.get(f"/api/projects/{pid}").json()
        bad = [(f["original_filename"], f["status"], f["error"]) for f in detail["files"] if f["status"] != "processed" and f["original_filename"] not in not_processed]
        assert not bad, bad
        assert {f["status"] for f in detail["files"] if f["original_filename"] in not_processed} == {"needs_ocr"}
        assert detail["stage"] == "review"
        ui = client.get(f"/api/projects/{pid}/report-data").json()
        v = {f["path"]: f["effective"] for s in ui["sections"] for f in s["fields"]}
        cells = {c["path"]: c["effective"] for s in ui["sections"] for t in s["tables"] for r in t["rows"] for c in r["cells"]}
        cells.update({c["path"]: c["effective"] for s in ui["sections"] for t in s["tables"] for c in t["totals"]})
        assert v["property.fields.name"] == "The Boardwalk" and v["property.fields.units"] == 338
        assert abs(cells["financials.tables.lines.rows.noi.ptd_actual"] - 636105.69) < 1
        assert abs(cells["financials.tables.lines.rows.total_revenue.ptd_actual"] - 1450137.81) < 1  # columns were swapped in the file:
        assert abs(cells["financials.tables.lines.rows.total_revenue.ptd_budget"] - 1603231) < 1     # header-driven mapping keeps them apart
        assert cells["financials.tables.lines.rows.noi.annual_budget"] is None  # the optional column was removed
        assert abs(cells["capex.tables.lines.totals.ptd_actual"] - 241077.37) < 1 and abs(cells["capex.tables.lines.totals.ptd_budget"] - 278896) < 1
        assert abs(cells["in_place_rent.tables.by_floor_plan.totals.current_rent"] - 1323.98) < 0.01
        assert abs(v["occupancy.fields.current_pct"] - 0.9053) < 1e-4 and v["occupancy.fields.change_bps"] == -118
        assert v["occupancy.fields.current_date"] == "2026-06-30"
        assert cells["occupancy.tables.new_leases.totals.count"] == 38
        assert abs(v["submarket.fields.vacancy"] - 0.1634) < 1e-3 and abs(cells["submarket.tables.comps.rows.the-ashlar.asking_rent"] - 1719) < 1
        trend = [t for s in ui["sections"] if s["key"] == "rent_trend" for t in s["tables"]][0]
        assert len(trend["rows"]) == 12, "the supplied rent chart workbook should drive the trend"
        assert v["rent_trend.fields.comp_name"] == "Comp Set"
        errors = [i["message"] for i in ui["issues"] if i["severity"] == "error"]
        assert len(errors) == 1 and "Lakeside Villas" in errors[0] and "set aside" in errors[0], errors
        assert any("Financials source: using" in i["message"] for i in ui["issues"]), "the duplicated export should be reported"
        assert any("no extractable text" in i["message"] for i in ui["issues"])
